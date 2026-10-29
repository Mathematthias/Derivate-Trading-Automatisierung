"""Thesen-Spur und TIW-Ereignisse (User-Entscheid 2026-10-29).

- Thesen-Karte: Pflichtfelder, Merge per thesen.update, kompakte Karte im Digest
- tiw.log: Validierung, Ablage in state/tiw_events.csv, Auswertung knapp/weit
- watchlist_meta reicht tiw_angekratzt durch
- state_export: Blaetter Thesen-Karten/Thesen-Audit nur, wenn es etwas zu zeigen gibt
"""
from __future__ import annotations

import json
import sys
from datetime import date, timedelta
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import inbox_apply  # noqa: E402
import state_yaml  # noqa: E402
import tiw_events  # noqa: E402

HEUTE = date(2026, 10, 29)

KARTE = {
    "konsens": "Markt bestraft AI-Capex",
    "abweichung": "Capex wird erhoeht, Zulieferer profitieren",
    "katalysator": [{"date": "2026-11-17", "ereignis": "NVDA Q3"}],
    "thesen_falsifikat": "Big-5 kuerzt Capex-Guidance",
    "preis_falsifikat": {"symbol": "^IXIC", "op": "<", "level": 25213.37, "text": "EMA200"},
    "horizont": "2026-12-02",
    "vehikel": ["VVSM.DE"],
}


def _state():
    return {"watchlist": {"entries": []}, "radar": {"rows": []}, "thesen": {"thesen": []}}


class TestKarte:
    def test_vollstaendig(self):
        assert state_yaml.karte_fehlt({"karte": KARTE}) == []

    def test_fehlende_felder(self):
        k = dict(KARTE, konsens="", vehikel=[])
        assert set(state_yaml.karte_fehlt({"karte": k})) == {"konsens", "vehikel"}

    def test_preis_falsifikat_ohne_op_zaehlt_als_fehlend(self):
        k = dict(KARTE, preis_falsifikat={"symbol": "^IXIC", "level": 1})
        assert state_yaml.karte_fehlt({"karte": k}) == ["preis_falsifikat"]

    def test_update_mischt_karte_feldweise(self):
        st = _state()
        state_yaml.apply_actions(st, [{"op": "thesen.update", "id": "TECH", "spur": "thesen",
                                       "budget_r": 2, "starter_anteil": 0.5, "karte": KARTE}], today=HEUTE)
        state_yaml.apply_actions(st, [{"op": "thesen.update", "id": "TECH",
                                       "karte": {"katalysator": [{"date": "2026-11-03", "ereignis": "AMD"}]}}],
                                 today=HEUTE)
        t = st["thesen"]["thesen"][0]
        assert t["karte"]["konsens"] == KARTE["konsens"]
        assert t["karte"]["katalysator"][0]["ereignis"] == "AMD"
        assert t["budget_r"] == 2 and t["spur"] == "thesen"

    def test_unbekanntes_kartenfeld_wird_verworfen(self):
        st = _state()
        log = state_yaml.apply_actions(st, [{"op": "thesen.update", "id": "X", "karte": {"foo": 1}}], today=HEUTE)
        assert any("verworfen" in l and "foo" in l for l in log)

    def test_digest_traegt_kompakte_karte(self):
        st = _state()
        state_yaml.apply_actions(st, [{"op": "thesen.update", "id": "TECH", "verdikt": "spielen",
                                       "spur": "thesen", "karte": KARTE}], today=HEUTE)
        t = state_yaml.thesen_aktiv(st["thesen"], HEUTE)[0]
        assert t["karte_fehlt"] == []
        assert t["katalysator"][0]["tage"] == 19
        assert t["horizont"] == "2026-12-02" and t["horizont_erreicht"] is False
        assert t["preis_falsifikat"]["op"] == "<"

    def test_alte_these_ohne_spur_meldet_keine_fehlende_karte(self):
        st = _state()
        state_yaml.apply_actions(st, [{"op": "thesen.update", "id": "OEL", "verdikt": "spielen"}], today=HEUTE)
        assert state_yaml.thesen_aktiv(st["thesen"], HEUTE)[0]["karte_fehlt"] is None

    def test_klasse_thesen_spur_bekannt(self):
        assert "thesen_spur" in state_yaml.KLASSEN


EV = {"date": "2026-10-29", "symbol": "CWC.DE", "leg": "A", "klasse": "breakout_retest",
      "richtung": "LONG", "tiw": 114.6, "schluss": 114.2, "atr": 2.18, "aktion": "angekratzt"}


class TestTiwLog:
    def test_op_validiert(self):
        st = _state()
        log = state_yaml.apply_actions(st, [{"op": "tiw.log", "event": dict(EV, aktion="weg")}], today=HEUTE)
        assert any("verworfen" in l for l in log)
        log = state_yaml.apply_actions(st, [{"op": "tiw.log", "event": {"symbol": "X"}}], today=HEUTE)
        assert any("verworfen" in l for l in log)
        assert not st.get("_tiw_events")

    def test_abstand_vorzeichen(self):
        assert tiw_events.abstand_atr("LONG", 114.6, 114.2, 2.0) == -0.2
        assert tiw_events.abstand_atr("SHORT", 50.0, 51.0, 2.0) == -0.5
        assert tiw_events.abstand_atr("SHORT", 50.0, 49.0, 2.0) == 0.5

    def test_inbox_schreibt_csv_idempotent(self, tmp_path):
        st = tmp_path / "state"
        st.mkdir()
        payload = {"schema": 1, "source": "morning-check", "actions": [{"op": "tiw.log", "event": EV}]}
        item = inbox_apply.InboxItem("k1", "INBOX-morning-check-x.json", payload)
        inbox_apply.apply_inbox(st, [item], today=HEUTE)
        rows = tiw_events.read_events(st / "tiw_events.csv")
        assert len(rows) == 1 and rows[0]["Aktion"] == "angekratzt"
        assert float(rows[0]["Abstand_ATR"]) == pytest.approx(-0.183, abs=1e-3)
        # gleiche Aktion aus einer zweiten Datei: keine Doppelzeile
        item2 = inbox_apply.InboxItem("k2", "INBOX-morning-check-y.json", payload)
        inbox_apply.apply_inbox(st, [item2], today=HEUTE)
        assert len(tiw_events.read_events(st / "tiw_events.csv")) == 1

    def test_fill_und_summary(self):
        rows = [tiw_events.event_to_row(EV),
                tiw_events.event_to_row(dict(EV, symbol="FRA.DE", richtung="SHORT", tiw=58.3,
                                             schluss=59.5, atr=1.5, aktion="archiviert"))]

        def closes(sym, start, end):
            base = 115.5 if sym == "CWC.DE" else 61.0
            return [(start + timedelta(days=i), base) for i in range(1, 15)]

        assert tiw_events.fill_outcomes(rows, closes, today=HEUTE + timedelta(days=20)) == 2
        assert rows[0]["Zurueck_10HT"] == "ja"      # Long wieder ueber der TIW
        assert rows[1]["Zurueck_10HT"] == "nein"    # Short weiter ueber der TIW
        s = tiw_events.summary(rows)
        assert s["knapp"]["n"] == 1 and s["knapp"]["quote"] == 1.0
        assert s["weit"]["n"] == 1 and s["weit"]["quote"] == 0.0
        assert s["urteil_moeglich"] is False
        assert "Urteil moeglich: nein" in tiw_events.render_md(rows, HEUTE)

    def test_fill_wartet_auf_genug_handelstage(self):
        rows = [tiw_events.event_to_row(EV)]
        n = tiw_events.fill_outcomes(rows, lambda s, a, b: [(a + timedelta(days=1), 115.0)], today=HEUTE)
        assert n == 0 and rows[0].get("Kurs_5HT") in (None, "")


def test_watchlist_meta_reicht_angekratzt_durch():
    block = {"entries": [{"symbol": "CWC.DE", "status": "aktiv", "klasse": "breakout_retest",
                          "legs": [{"label": "A", "gate": "🔴", "text": "TIW = 1D-Schluss unter 114,60€"}],
                          "tiw_angekratzt": {"date": "2026-10-29", "abstand_atr": -0.18}}]}
    m = state_yaml.watchlist_meta(block)[0]
    assert m["tiw_angekratzt"]["date"] == "2026-10-29"


class TestExport:
    def test_ohne_thesen_spur_keine_neuen_blaetter(self, tmp_path):
        import state_export as se
        from openpyxl import Workbook
        wb = Workbook()
        counts = se.export_thesen(wb, _state(), tmp_path)
        assert counts == {}
        assert "Thesen-Karten" not in wb.sheetnames

    def test_karte_und_audit_im_journal(self, tmp_path):
        import state_export as se
        from openpyxl import Workbook
        st = _state()
        state_yaml.apply_actions(st, [{"op": "thesen.update", "id": "TECH", "titel": "AI-Zulieferer",
                                       "verdikt": "spielen", "spur": "thesen", "budget_r": 2, "karte": KARTE}],
                                 today=HEUTE)
        (tmp_path / "thesen_audit.csv").write_text("These,Start,R_gesamt\nTECH,2026-10-30,\n", encoding="utf-8")
        wb = Workbook()
        counts = se.export_thesen(wb, st, tmp_path)
        assert counts["karten"] == 1 and counts["audit"] == 1
        ws = wb["Thesen-Karten"]
        assert ws.cell(2, 1).value == "TECH"
        assert "^IXIC Schluss < 25213.37" in ws.cell(2, 10).value
        assert ws.cell(2, 15).value == "—"
        assert wb["Thesen-Audit"].cell(2, 1).value == "TECH"
