"""state_export (YAML -> Journal-Blaetter) und mfe_mae (Excursion in R)."""
import os
import sys
from datetime import date

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

_HERE = os.path.dirname(os.path.abspath(__file__))
_SRC = os.path.join(os.path.dirname(_HERE), "src")
sys.path.insert(0, _SRC)

import mfe_mae  # noqa: E402
import state_export as se  # noqa: E402
import state_yaml as sy  # noqa: E402


# ---------------------------------------------------------------------------
# state_export
# ---------------------------------------------------------------------------

def _journal(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Watchlist"
    ws.append(se.WL_HEADER)
    ws.append(["Alt AG", "ALT.DE", "LONG", "🟡", "alter Trigger", None, None, None, None, "x", "2026-01-01", "2026-02-01"])
    ws2 = wb.create_sheet("Watchlist-Archiv")
    ws2.append(se.ARCHIV_HEADER)
    ws2.append(["Alt2", "ALT2.DE", "LONG", "A) x", "y", "📦 ARCHIVIERT 2026-05-01 — alt", "2026-04-01", "2026-04-15"])
    ws3 = wb.create_sheet("Termin-Radar")
    ws3.append(["TERMIN-RADAR — Titel"])
    ws3.append(["Kat: 1=..."])
    ws3.append([])
    ws3.append(se.RADAR_HEADER)
    ws3.append(["2026-01-01", 1, "alt", "", "", "offen", ""])
    ws4 = wb.create_sheet("Geschlossene Trades")
    ws4.append(["Nr.", "G/V"])
    ws4.append([1, 10])
    ws4.append(["Summe", "=SUM(B2:B2)"])
    wb.save(path)


def _state():
    return {
        "watchlist": {"updated": "2026-10-04T20:00:00+02:00", "entries": [
            {"symbol": "SIE.DE", "name": "Siemens", "direction": "LONG", "klasse": "trend_pullback", "anker": "EMA20",
             "status": "aktiv", "added": "2026-09-29", "expiry": "2026-10-19", "journal_row": 3, "treiber": "Q4",
             "legs": [{"label": "A", "gate": "🟡", "text": "TPL"}, {"label": "C", "gate": "🔴", "text": "Alt"}]},
            {"symbol": "JST.DE", "name": "JOST", "direction": "LONG", "klasse": "position_monitor", "status": "position",
             "trade_nr": "100", "added": "2026-09-02", "legs": [{"label": "A", "gate": "📍", "text": "Position #100"}]},
            {"symbol": "ZTS", "name": "Zoetis", "direction": "SHORT", "klasse": "trend_pullback", "status": "archiviert",
             "archived": {"date": "2026-10-04", "reason": "ohne Katalysator"}, "added": "2026-10-02",
             "legs": [{"label": "A", "gate": "🟡", "text": "STPL"}]},
            {"symbol": "ALT2.DE", "name": "Alt2", "direction": "LONG", "klasse": "sonstige", "status": "archiviert",
             "archived": {"date": "2026-05-01", "reason": "alt"}, "legs": [{"label": "A", "gate": "🔴", "text": "x"}]},
        ]},
        "radar": {"updated": "2026-10-04T20:00:00+02:00", "rows": [
            {"id": "Z2", "date": "2026-11-01", "kat": 2, "ereignis": "spaeter", "werte": "", "wirkung": "", "status": "offen", "quelle": "q"},
            {"id": "Z1", "date": "2026-10-10", "kat": 1, "ereignis": "frueher", "werte": "A", "wirkung": "w", "status": "offen", "quelle": "q", "note": "KORR"},
        ]},
        "thesen": {"thesen": []},
    }


class TestStateExport:

    def test_blaetter_werden_aus_yaml_geschrieben(self, tmp_path):
        src, dst = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
        _journal(src)
        counts = se.export_journal(_state(), src, dst)
        assert counts == {"watchlist": 2, "archiv_neu": 1, "radar": 2}
        wb = load_workbook(dst)
        ws = wb["Watchlist"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert [r[1] for r in rows] == ["JST.DE", "SIE.DE"]          # Position zuerst, Archiv weg
        assert rows[0][3] == "📍" and "Position #100" in rows[0][2]
        assert rows[1][3] == "🟡" and rows[1][4] == "TPL"
        assert rows[1][7] == "🔴" and rows[1][8] == "Alt"            # Leg C in Spalte H/I
        assert rows[1][9].startswith("Treiber: Q4")
        # Gate-Validierung umfasst alle Datenzeilen
        dv = ws.data_validations.dataValidation[0]
        assert "D2:D3" in str(dv.sqref)

    def test_archiv_wird_angehaengt_ohne_duplikat(self, tmp_path):
        src, dst = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
        _journal(src)
        se.export_journal(_state(), src, dst)
        ws = load_workbook(dst)["Watchlist-Archiv"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert [r[1] for r in rows] == ["ALT2.DE", "ZTS"]            # ALT2 nur einmal
        assert "ARCHIVIERT 2026-10-04" in rows[1][5] and "ohne Katalysator" in rows[1][5]

    def test_radar_sortiert_und_mit_note(self, tmp_path):
        src, dst = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
        _journal(src)
        se.export_journal(_state(), src, dst)
        ws = load_workbook(dst)["Termin-Radar"]
        assert ws.cell(row=4, column=1).value == "Datum"           # Kopf bleibt
        rows = list(ws.iter_rows(min_row=5, values_only=True))
        assert [r[0] for r in rows] == ["2026-10-10", "2026-11-01"]
        assert rows[0][6] == "q | KORR"

    def test_andere_blaetter_und_formeln_unangetastet(self, tmp_path):
        src, dst = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
        _journal(src)
        se.export_journal(_state(), src, dst)
        ws = load_workbook(dst)["Geschlossene Trades"]
        assert ws["B3"].value == "=SUM(B2:B2)"
        assert "State-Export" in load_workbook(dst).sheetnames

    def test_export_ist_idempotent(self, tmp_path):
        src, mid, dst = tmp_path / "in.xlsx", tmp_path / "mid.xlsx", tmp_path / "out.xlsx"
        _journal(src)
        se.export_journal(_state(), src, mid)
        counts = se.export_journal(_state(), mid, dst)
        assert counts["archiv_neu"] == 0 and counts["watchlist"] == 2
        a = list(load_workbook(mid)["Watchlist"].iter_rows(values_only=True))
        b = list(load_workbook(dst)["Watchlist"].iter_rows(values_only=True))
        assert a == b

    def test_cli(self, tmp_path):
        src, dst = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
        _journal(src)
        sy.save_state(tmp_path / "state", _state())
        assert se.main(["--journal", str(src), "--out", str(dst), "--state-dir", str(tmp_path / "state")]) == 0
        assert dst.exists()


# ---------------------------------------------------------------------------
# mfe_mae
# ---------------------------------------------------------------------------

def _bars(highs, lows, start="2026-09-01"):
    idx = pd.bdate_range(start, periods=len(highs))
    return pd.DataFrame({"High": highs, "Low": lows}, index=idx)


class TestComputeMfeMae:

    def test_long_ausstiegsfehler(self):
        # Entry 100, SL 98 (Risiko 2). Hoch 103 (= +1,5 R), Tief 99 (= -0,5 R)
        df = _bars([101, 103, 102, 101], [99.5, 101, 99, 100])
        ex = mfe_mae.compute_mfe_mae(df, "Long", 100, 98, date(2026, 9, 1), date(2026, 9, 4))
        assert ex.mfe_r == 1.5 and ex.mae_r == 0.5 and ex.bars == 4
        assert mfe_mae.classify(ex.mfe_r, ex.mae_r, -1.0) == "Ausstiegsfehler (lief, dann weg)"

    def test_long_einstiegsfehler(self):
        df = _bars([100.4, 100.2, 99.0], [99.0, 98.3, 97.9])
        ex = mfe_mae.compute_mfe_mae(df, "LONG", 100, 98, date(2026, 9, 1), date(2026, 9, 3))
        assert ex.mfe_r == 0.2 and ex.mae_r == 1.05
        assert mfe_mae.classify(ex.mfe_r, ex.mae_r, -1.0) == "Einstiegsfehler (lief nie)"

    def test_short_spiegel(self):
        # Short Entry 100, SL 102 (Risiko 2). Tief 97 (= +1,5 R), Hoch 101 (= -0,5 R)
        df = _bars([100.5, 101, 99], [99, 98, 97])
        ex = mfe_mae.compute_mfe_mae(df, "Short", 100, 102, date(2026, 9, 1), date(2026, 9, 3))
        assert ex.mfe_r == 1.5 and ex.mae_r == 0.5

    def test_fenster_beschraenkt_auf_haltedauer(self):
        df = _bars([101, 110, 101, 101], [99, 99, 99, 90])
        ex = mfe_mae.compute_mfe_mae(df, "Long", 100, 98, date(2026, 9, 1), date(2026, 9, 1))
        assert ex.mfe_r == 0.5 and ex.bars == 1

    def test_risiko_null_oder_leer_gibt_none(self):
        df = _bars([101], [99])
        assert mfe_mae.compute_mfe_mae(df, "Long", 100, 100, date(2026, 9, 1), date(2026, 9, 1)) is None
        assert mfe_mae.compute_mfe_mae(df, "Long", 100, 98, date(2026, 10, 1), date(2026, 10, 2)) is None

    def test_gewinner_klassifikation(self):
        assert mfe_mae.classify(2.3, 0.3, 1.8) == "Gewinner (lief)"
        assert mfe_mae.classify(0.6, 0.3, 0.2) == "Gewinner (knapp)"
        assert mfe_mae.classify(0.3, 0.9, -1.0) == "Einstiegsfehler (lief nie)"
        assert mfe_mae.classify(0.7, 0.9, -1.0) == "Stop im Rauschen?"


class TestParseDate:

    @pytest.mark.parametrize("raw,exp", [
        ("2026-03-02", date(2026, 3, 2)),
        ("23.03.2026", date(2026, 3, 23)),
        ("11.03./13.03.2026", date(2026, 3, 11)),
        (date(2026, 1, 5), date(2026, 1, 5)),
        ("", None), (None, None), ("bald", None),
    ])
    def test_formate(self, raw, exp):
        assert mfe_mae.parse_date(raw) == exp

    def test_zahlen_deutsch_und_englisch(self):
        assert mfe_mae._num("3.481,00€") == 3481.0
        assert mfe_mae._num("651.77$") == 651.77
        assert mfe_mae._num("98,5") == 98.5
        assert mfe_mae._num(None) is None


class TestRunOffline:

    def _audit(self, path):
        wb = Workbook()
        ws = wb.active
        ws.title = mfe_mae.AUDIT_SHEET
        ws.append(["Nr", "Symbol", "Richtung", "Kauf", "Verkauf", "EntryU", "SL_U", "R", "Klasse"])
        ws.append([101, "JST.DE", "Long", "2026-09-01", "2026-09-04", 100, 98, -1.0, "TPL"])
        ws.append([102, "", "Long", "2026-09-01", "2026-09-04", 100, 98, None, "TPL"])      # kein Symbol
        ws.append([103, "XXX", "Short", "2026-09-01", "2026-09-03", 100, 102, 1.2, "G"])
        ws.append([104, "NIX", "Long", "2026-09-01", "2026-09-03", 100, 98, None, ""])       # keine Daten
        wb.save(path)

    def test_run_schreibt_spalten_und_csv(self, tmp_path):
        src, dst, csv = tmp_path / "a.xlsx", tmp_path / "b.xlsx", tmp_path / "m.csv"
        self._audit(src)

        def fake_fetch(symbol, start, end):
            if symbol == "NIX":
                return None
            if symbol == "JST.DE":
                return _bars([101, 103, 102, 101], [99.5, 101, 99, 100])
            return _bars([100.5, 101, 99], [99, 98, 97])

        table, summary = mfe_mae.run(src, dst, csv, fetch=fake_fetch)
        ok = {t["nr"]: t for t in table if t["status"] == "ok"}
        assert set(ok) == {"101", "103"}
        assert ok["101"]["klasse"] == "Ausstiegsfehler (lief, dann weg)"
        assert ok["103"]["mfe_r"] == 1.5
        ws = load_workbook(dst)[mfe_mae.AUDIT_SHEET]
        hdr = [c.value for c in ws[1]]
        assert hdr[-3:] == ["MFE_R", "MAE_R", "Bars"]
        assert ws.cell(row=2, column=hdr.index("MFE_R") + 1).value == 1.5
        assert ws.cell(row=3, column=hdr.index("MFE_R") + 1).value is None
        assert csv.exists() and "uebersprungen (Symbol)" in csv.read_text(encoding="utf-8")
        assert "2 von 4 Zeilen gerechnet" in summary and "Nr 104 NIX: keine Kursdaten" in summary


def test_archiv_ohne_datum_wird_nicht_doppelt_angehaengt(tmp_path):
    """Regression AMZN 2026-10-05: archived.date leer, Datum nur im Grund-Text."""
    j = tmp_path / "in.xlsx"
    _journal(j)
    state = {
        "watchlist": {"entries": [{
            "symbol": "AMZN", "name": "Amazon", "direction": "LONG", "status": "archiviert",
            "legs": [{"label": "A", "gate": "🔴", "text": "x"}],
            "archived": {"date": "", "reason": "Position geschlossen 2026-09-16"},
        }]},
        "radar": {"rows": []}, "thesen": {"thesen": []},
    }
    out1, out2 = tmp_path / "o1.xlsx", tmp_path / "o2.xlsx"
    c1 = se.export_journal(state, j, out1)
    c2 = se.export_journal(state, out1, out2)
    assert c1["archiv_neu"] == 1 and c2["archiv_neu"] == 0
    ws = load_workbook(out2)["Watchlist-Archiv"]
    assert sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r[1] == "AMZN") == 1


# ---------------------------------------------------------------------------
# CSV-Modus und Merge (Fix 2026-10-08)
# ---------------------------------------------------------------------------

def _bars_csv():
    idx = pd.to_datetime(["2026-09-01", "2026-09-02", "2026-09-03", "2026-09-04"])
    return pd.DataFrame({"High": [101.0, 104.0, 102.0, 100.5], "Low": [99.5, 100.5, 99.0, 97.5]}, index=idx)


def test_csv_mode_and_merge(tmp_path):
    inp = tmp_path / "trade_audit.csv"
    pd.DataFrame([
        {"TradeID": "D-091", "Symbol": "FRE.DE", "Richtung": "Long", "Kauf": "2026-09-01", "Verkauf": "2026-09-04",
         "EntryU": "100", "SL_U": "98", "R": "-1.21"},
        {"TradeID": "A-010", "Symbol": "LLY", "Richtung": "Long", "Kauf": "2026-09-01", "Verkauf": "",
         "EntryU": "100", "SL_U": "98", "R": ""},
        {"TradeID": "D-065", "Symbol": "CL=F", "Richtung": "Long", "Kauf": "2026-09-01", "Verkauf": "2026-09-02",
         "EntryU": "", "SL_U": "", "R": "-1.02"},
    ]).to_csv(inp, index=False)
    out = tmp_path / "mfe_mae.csv"
    table, summary = mfe_mae.run_csv(inp, out, fetch=lambda s, a, b: _bars_csv(), today=date(2026, 9, 4))
    by = {t["TradeID"]: t for t in table}
    assert by["D-091"]["MFE_R"] == 2.0 and by["D-091"]["MAE_R"] == 1.25
    assert by["D-091"]["MFE_Datum"] == "2026-09-02" and by["D-091"]["MAE_Datum"] == "2026-09-04"
    assert by["D-091"]["Klasse"] == "Ausstiegsfehler (lief, dann weg)"
    assert by["A-010"]["Status"] == "ok (offen)" and by["A-010"]["Klasse"] == "offen"
    assert by["D-065"]["Status"].startswith("uebersprungen")
    # Merge ins Journal: D-091 vorhanden (Nr int), A-010 wird angehaengt
    j = tmp_path / "j.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "Trade-Audit"
    ws.append(["Nr", "Instrument", "Symbol", "Richtung", "Kauf", "Verkauf", "R", "EntryU", "SL_U", "Status", "MFE_R", "MAE_R", "Bars"])
    ws.append([91, "FRE KO", None, "Long", "2026-09-01", "2026-09-04", -1.21, None, None, "CLOSED", None, None, None])
    wb.save(j)
    st = mfe_mae.merge_into_journal(j, out, tmp_path / "j2.xlsx", inp)
    assert st == {"gefuellt": 2, "angehaengt": 1}
    ws2 = load_workbook(tmp_path / "j2.xlsx")["Trade-Audit"]
    assert ws2["C2"].value == "FRE.DE" and ws2["H2"].value == 100 and ws2["K2"].value == 2.0
    assert ws2["A3"].value == "A-010" and ws2["J3"].value == "OFFEN"


# ---------------------------------------------------------------------------
# Gegenfakten (A) und Einstiegs-Kontext (C), 2026-10-08
# ---------------------------------------------------------------------------

def _ohlc(rows, start="2026-08-03"):
    idx = pd.bdate_range(start, periods=len(rows))
    return pd.DataFrame(rows, columns=["Open", "High", "Low", "Close"], index=idx)


def test_cf_be_rettet_ruecklaeufer():
    # Long, Entry 100, SL 98 (R=2). Tag1 Hoch 101,6 (+0,8R), Tag2 faellt auf 97 (SL).
    df = _ohlc([[100, 100.5, 99.5, 100], [100.2, 101.6, 100.1, 101], [100.5, 100.6, 97.0, 97.5]])
    d = df.index.date
    cf = mfe_mae.counterfactuals(df, "Long", 100, 98, d[0], d[2], -1.0)
    assert cf["CF_BE07"] == 0.0          # +0,7R an Tag 1 erreicht, ab Tag 2 Stop auf Einstand
    assert cf["CF_BE10"] == -1.0         # +1,0R (102) nie erreicht
    assert cf["CF_P10"] == -1.0          # 1/3-Teilverkauf nie ausgeloest


def test_cf_gap_und_teilverkauf():
    # Long, Entry 100, SL 98. Tag1 Hoch 102,2 (+1,1R), Tag2 oeffnet mit Gap bei 99 -> -0,5R
    df = _ohlc([[100, 100.4, 99.8, 100.2], [100.3, 102.2, 100.2, 102], [99.0, 99.2, 97.5, 97.8]])
    d = df.index.date
    cf = mfe_mae.counterfactuals(df, "Long", 100, 98, d[0], d[2], -1.05)
    assert cf["CF_BE07"] == -0.5 and cf["CF_BE10"] == -0.5
    assert cf["CF_P10"] == round(1/3 + 2/3 * -1.05, 2)
    assert "Gap" in cf["CF_Hinweis"]


def test_cf_short_gespiegelt():
    # Short, Entry 100, SL 102 (R=2). Tag1 Tief 98,4 (+0,8R), Tag2 steigt auf 103
    df = _ohlc([[100, 100.5, 99.6, 99.8], [99.7, 99.8, 98.4, 99], [99.5, 103, 99.4, 102.8]])
    d = df.index.date
    cf = mfe_mae.counterfactuals(df, "Short", 100, 102, d[0], d[2], -1.0)
    assert cf["CF_BE07"] == 0.0


def test_entry_context_lage_und_gap():
    rows = [[100, 101, 99, 100]] * 20 + [[103, 104, 102, 103.5], [101, 101.5, 99, 100]]
    df = _ohlc(rows)
    d = df.index.date
    ctx = mfe_mae.entry_context(df, "Long", 104, d[20], d[21])
    assert ctx["Einstieg_Lage_Pct"] == 100          # Kauf am Tageshoch
    assert ctx["Gap_Einstieg_ATR"] > 1              # Eroeffnung 103 vs Vortag 100 bei ATR ~2
    assert ctx["Max_Gap_gegen_ATR"] > 0.5           # Folgetag oeffnet 2,5 unter Schluss 103,5
