#!/usr/bin/env python3
"""
state_export.py — state/*.yaml -> Journal-xlsx (Watchlist, Watchlist-Archiv,
Termin-Radar). Workflow B, System-Review 2026-10-04.

Das Journal ist ab jetzt fuer diese drei Blaetter ABGELEITET: Quelle ist
``state/watchlist.yaml`` und ``state/radar.yaml`` im Repo. Dieses Skript nimmt
ein bestehendes Journal, leert die drei Blaetter unterhalb der Kopfzeile(n)
und schreibt den Repo-Zustand hinein. Alle anderen Blaetter (Geschlossene
Trades, Steuer-Blaetter, Lektionen ...) bleiben byte-fuer-byte-inhaltlich
unangetastet — openpyxl laedt OHNE data_only, Formeln bleiben Formeln.

Regeln:
  * Blatt "Watchlist": nur status in (position, aktiv). Positionen zuerst
    (📍 im Gate A), dann aktive Zeilen in Journal-Reihenfolge (journal_row),
    neue Eintraege (ohne journal_row) nach Hinzufuege-Datum.
  * Blatt "Watchlist-Archiv": archivierte Eintraege werden angehaengt, sofern
    (Symbol, Archiv-Datum) dort noch nicht steht. Bestehende Archiv-Zeilen
    bleiben, wie sie sind.
  * Blatt "Termin-Radar": Kopf (Titel, Legende, Header) bleibt, Datenzeilen
    werden komplett aus radar.yaml neu geschrieben, sortiert nach Datum.
  * Datenvalidierung der Gate-Spalten (🟢,🟡,⏳,🔴) wird auf die neue
    Zeilenzahl ausgedehnt.

Aufruf:
    python src/state_export.py --journal Trading_Journal_20261002c.xlsx \
        --out Trading_Journal_20261004b.xlsx --state-dir ./state
    python src/state_export.py --journal in.xlsx --out out.xlsx --upload   # + Drive
"""
from __future__ import annotations

import argparse
import copy
import logging
import os
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

import state_yaml

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("state_export")

KLASSE_LABEL = {
    "trend_pullback": "Trend-Pullback",
    "grinder": "Grinder-Continuation",
    "breakout_retest": "Breakout-Retest",
    "thesen_korb": "Thesen-Korb",
    "insider": "Insider-Cluster",
    "news_catalyst": "News-Katalysator",
    "event_watch": "Event-Watch",
    "position_monitor": "Position",
    "sonstige": "Sonstige",
}

WL_HEADER = ["Aktie", "Symbol", "Richtung", "🚦 A", "Trigger A", "🚦 B", "Trigger B",
             "🚦 C", "Trigger C", "Bemerkungen", "Datum hinzugefügt", "Verfallsdatum"]
ARCHIV_HEADER = ["Aktie", "Symbol", "Richtung", "Entry-Trigger", "These", "Status",
                 "Datum hinzugefügt", "Verfallsdatum"]
RADAR_HEADER = ["Datum", "Kat", "Ereignis", "Werte (Ticker)", "Wirkrichtung", "Status", "Quelle / Abruf"]


# ---------------------------------------------------------------------------
# Hilfen
# ---------------------------------------------------------------------------

def _row_style(ws, row: int, ncols: int) -> list[dict]:
    """Kopiert die Zellformate einer Zeile (fuer gleichmaessige neue Zeilen)."""
    out = []
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        out.append({
            "font": copy.copy(cell.font), "alignment": copy.copy(cell.alignment),
            "border": copy.copy(cell.border), "fill": copy.copy(cell.fill),
            "number_format": cell.number_format,
        })
    return out


def _apply_style(ws, row: int, styles: list[dict]) -> None:
    for c, st in enumerate(styles, start=1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.alignment = st["font"], st["alignment"]
        cell.border, cell.fill = st["border"], st["fill"]
        cell.number_format = st["number_format"]


def _clear_rows(ws, first_row: int) -> None:
    if ws.max_row >= first_row:
        ws.delete_rows(first_row, ws.max_row - first_row + 1)


def _header_row(ws, first_col_text: str, scan: int = 10) -> int:
    for r in range(1, scan + 1):
        if str(ws.cell(row=r, column=1).value or "").strip().lower() == first_col_text.lower():
            return r
    raise ValueError(f"Kopfzeile '{first_col_text}' nicht in den ersten {scan} Zeilen von {ws.title}")


def _richtung_text(e: dict) -> str:
    direction = str(e.get("direction", "LONG")).upper()
    klasse = KLASSE_LABEL.get(e.get("klasse"), e.get("klasse") or "")
    anker = e.get("anker")
    if e.get("status") == "position":
        return f"{direction} 📍 Position{(' #' + str(e['trade_nr'])) if e.get('trade_nr') else ''}"
    parts = [klasse] if klasse else []
    if anker:
        parts.append(f"Anker {anker}")
    return f"{direction} ({', '.join(parts)})" if parts else direction


def _legs_text(e: dict) -> str:
    return " · ".join(
        f"{l.get('gate', '')} {l.get('label', '')}) {l.get('text', '')}".strip()
        for l in (e.get("legs") or [])
    )


def _sort_key_active(e: dict):
    pos = 0 if e.get("status") == "position" else 1
    jr = e.get("journal_row")
    return (pos, 0 if jr else 1, jr or 0, str(e.get("added") or ""), e.get("symbol", ""))


# ---------------------------------------------------------------------------
# Watchlist
# ---------------------------------------------------------------------------

def export_watchlist(ws, entries: list[dict]) -> int:
    hdr = _header_row(ws, "Aktie")
    ncols = len(WL_HEADER)
    styles = _row_style(ws, hdr + 1, ncols) if ws.max_row > hdr else None
    _clear_rows(ws, hdr + 1)

    live = sorted((e for e in entries if e.get("status") in ("aktiv", "position")), key=_sort_key_active)
    r = hdr
    for e in live:
        r += 1
        legs = {str(l.get("label", "")).upper(): l for l in (e.get("legs") or [])}
        gate_a = "📍" if e.get("status") == "position" else (legs.get("A", {}).get("gate") or "")
        vals = [
            e.get("name") or e.get("symbol"), e.get("symbol"), _richtung_text(e),
            gate_a, legs.get("A", {}).get("text", ""),
            legs.get("B", {}).get("gate", "") if "B" in legs else None, legs.get("B", {}).get("text") if "B" in legs else None,
            legs.get("C", {}).get("gate", "") if "C" in legs else None, legs.get("C", {}).get("text") if "C" in legs else None,
            _bemerkung(e), e.get("added") or None, e.get("expiry") or None,
        ]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
        if styles:
            _apply_style(ws, r, styles)
        e["_export_row"] = r

    # Gate-Validierung auf neue Laenge ziehen
    last = max(r, hdr + 1)
    ws.data_validations.dataValidation = []
    dv = DataValidation(type="list", formula1='"🟢,🟡,⏳,🔴,📍"', allow_blank=True)
    dv.add(f"D{hdr + 1}:D{last}")
    dv.add(f"F{hdr + 1}:F{last}")
    dv.add(f"H{hdr + 1}:H{last}")
    ws.add_data_validation(dv)
    return len(live)


def _bemerkung(e: dict) -> str:
    bits = []
    if e.get("treiber"):
        bits.append(f"Treiber: {e['treiber']}")
    if e.get("notes"):
        bits.append(str(e["notes"]))
    return " | ".join(bits)[:1000]


# ---------------------------------------------------------------------------
# Watchlist-Archiv
# ---------------------------------------------------------------------------

def export_archive(ws, entries: list[dict]) -> int:
    hdr = _header_row(ws, "Aktie")
    ncols = len(ARCHIV_HEADER)
    styles = _row_style(ws, hdr + 1, ncols) if ws.max_row > hdr else None
    existing = set()
    for row in ws.iter_rows(min_row=hdr + 1, values_only=True):
        sym = str(row[1] or "").strip().upper()
        status = str(row[5] or "")
        if sym:
            existing.add((sym, _first_iso(status)))
    added = 0
    r = ws.max_row
    for e in sorted((e for e in entries if e.get("status") == "archiviert"),
                    key=lambda x: (str((x.get("archived") or {}).get("date") or ""), x.get("symbol", ""))):
        arch = e.get("archived") or {}
        status_txt = f"📦 ARCHIVIERT {arch.get('date') or '?'} — {arch.get('reason') or ''}".strip(" —")
        # Schluessel genau so bilden wie beim Lesen des Blatts (_first_iso auf dem
        # Status-Text). Vorher: Datum aus archived.date — fehlte es und stand ein
        # Datum nur im Grund-Text, passten Lese- und Schreibschluessel nie zusammen
        # und die Zeile wurde bei JEDEM Export erneut angehaengt (AMZN, 2026-10-05).
        key = (str(e.get("symbol", "")).upper(), _first_iso(status_txt))
        if key in existing:
            continue
        r += 1
        vals = [
            e.get("name") or e.get("symbol"), e.get("symbol"), _richtung_text({**e, "status": "aktiv"}),
            _legs_text(e), _bemerkung(e), status_txt, e.get("added") or None, e.get("expiry") or None,
        ]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
        if styles:
            _apply_style(ws, r, styles)
        existing.add(key)
        added += 1
    return added


def _first_iso(text: str) -> str:
    import re
    m = re.search(r"\d{4}-\d{2}-\d{2}", text or "")
    return m.group(0) if m else ""


# ---------------------------------------------------------------------------
# Termin-Radar
# ---------------------------------------------------------------------------

def export_radar(ws, rows: list[dict]) -> int:
    hdr = _header_row(ws, "Datum")
    ncols = len(RADAR_HEADER)
    styles = _row_style(ws, hdr + 1, ncols) if ws.max_row > hdr else None
    _clear_rows(ws, hdr + 1)
    r = hdr
    for row in sorted(rows, key=lambda x: (str(x.get("date") or ""), str(x.get("id") or ""))):
        r += 1
        quelle = str(row.get("quelle") or "")
        if row.get("note"):
            quelle = f"{quelle} | {row['note']}".strip(" |")
        vals = [row.get("date"), row.get("kat"), row.get("ereignis"), row.get("werte"),
                row.get("wirkung"), row.get("status", "offen"), quelle]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
        if styles:
            _apply_style(ws, r, styles)
    return r - hdr


# ---------------------------------------------------------------------------
# Info-Blatt
# ---------------------------------------------------------------------------

def export_info(wb, state: dict, counts: dict) -> None:
    name = "State-Export"
    ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
    _clear_rows(ws, 1)
    lines = [
        ("Blatt", "Quelle", "Stand (YAML updated)", "Zeilen"),
        ("Watchlist", "state/watchlist.yaml (status aktiv/position)", state["watchlist"].get("updated"), counts["watchlist"]),
        ("Watchlist-Archiv", "state/watchlist.yaml (status archiviert, neu angehaengt)", state["watchlist"].get("updated"), counts["archiv_neu"]),
        ("Termin-Radar", "state/radar.yaml", state["radar"].get("updated"), counts["radar"]),
        ("", "", "", ""),
        ("Exportiert", datetime.now().astimezone().isoformat(timespec="seconds"), "", ""),
        ("Hinweis", "Diese drei Blaetter sind ABGELEITET. Aenderungen gehoeren in state/*.yaml (Commit) oder als INBOX-Aktion, nicht hierher.", "", ""),
    ]
    for r, vals in enumerate(lines, start=1):
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 8


# ---------------------------------------------------------------------------
# Orchestrierung
# ---------------------------------------------------------------------------

def export_journal(state: dict, journal_in: Path, journal_out: Path) -> dict:
    wb = load_workbook(journal_in)          # NICHT data_only: Formeln bleiben
    entries = state["watchlist"].get("entries", []) or []
    counts = {"watchlist": 0, "archiv_neu": 0, "radar": 0}
    if "Watchlist" in wb.sheetnames:
        counts["watchlist"] = export_watchlist(wb["Watchlist"], entries)
    else:
        raise ValueError("Blatt 'Watchlist' fehlt im Journal")
    if "Watchlist-Archiv" in wb.sheetnames:
        counts["archiv_neu"] = export_archive(wb["Watchlist-Archiv"], entries)
    if "Termin-Radar" in wb.sheetnames:
        counts["radar"] = export_radar(wb["Termin-Radar"], state["radar"].get("rows", []) or [])
    export_info(wb, state, counts)
    journal_out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(journal_out)
    logger.info(f"Export: {counts} -> {journal_out}")
    return counts


def main(argv: Optional[list[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--journal", required=True, help="Eingangs-Journal (.xlsx)")
    ap.add_argument("--out", required=True, help="Ausgabe-Journal (.xlsx)")
    ap.add_argument("--state-dir", default=os.environ.get("STATE_DIR", "./state"))
    ap.add_argument("--upload", action="store_true", help="Ergebnis zusaetzlich nach BRIEFING_FOLDER_ID hochladen")
    args = ap.parse_args(argv)

    state = state_yaml.load_state(args.state_dir)
    counts = export_journal(state, Path(args.journal), Path(args.out))
    print(f"Watchlist {counts['watchlist']} Zeilen, Archiv +{counts['archiv_neu']}, Radar {counts['radar']} Zeilen -> {args.out}")

    if args.upload:
        from drive_writer import build_drive_service, write_binary_file
        folder = os.environ.get("BRIEFING_FOLDER_ID")
        if not folder:
            logger.error("--upload braucht BRIEFING_FOLDER_ID")
            return 2
        svc = build_drive_service()
        write_binary_file(
            svc, folder, Path(args.out).name, Path(args.out).read_bytes(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
